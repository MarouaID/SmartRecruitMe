"""
Script pour générer un CV parfaitement aligné avec l'offre Full Stack
"""
from docx import Document

def create_perfect_match_cv():
    """CV - Candidat parfait pour Développeur Full Stack"""
    doc = Document()
    
    doc.add_heading('Sarah ALAMI', 0)
    doc.add_paragraph('Développeuse Full Stack Senior')
    doc.add_paragraph('Email: sarah.alami@email.com | Tél: +212 6 77 88 99 00')
    doc.add_paragraph('GitHub: github.com/sarahalami | LinkedIn: linkedin.com/in/sarahalami')
    doc.add_paragraph('Localisation: Casablanca, Maroc')
    
    doc.add_heading('Profil', 1)
    doc.add_paragraph(
        'Développeuse Full Stack Senior avec 3 ans d\'expérience spécialisée en React, Node.js, '
        'Python, MongoDB et Docker. Passionnée par la création d\'applications web performantes '
        'et scalables avec des architectures modernes et des pratiques DevOps.'
    )
    
    doc.add_heading('Expérience Professionnelle', 1)
    
    doc.add_heading('Développeuse Full Stack Senior - InnovateTech', 2)
    doc.add_paragraph('Janvier 2022 - Présent | Casablanca, Maroc')
    doc.add_paragraph('• Développement d\'applications web avec React et TypeScript')
    doc.add_paragraph('• Création d\'APIs RESTful robustes avec Node.js et Express')
    doc.add_paragraph('• Développement de microservices en Python avec FastAPI')
    doc.add_paragraph('• Gestion de bases de données MongoDB avec optimisation des requêtes')
    doc.add_paragraph('• Containerisation complète des applications avec Docker et Docker Compose')
    doc.add_paragraph('• Mise en place de pipelines CI/CD avec Docker')
    doc.add_paragraph('• Architecture et déploiement d\'applications conteneurisées')
    
    doc.add_heading('Développeuse Full Stack - WebSolutions', 2)
    doc.add_paragraph('Mars 2021 - Décembre 2021 | Rabat, Maroc')
    doc.add_paragraph('• Développement frontend avec React, Redux et Material-UI')
    doc.add_paragraph('• Backend Node.js avec Express et MongoDB')
    doc.add_paragraph('• Scripts Python pour automatisation et traitement de données')
    doc.add_paragraph('• Déploiement d\'applications avec Docker')
    doc.add_paragraph('• Intégration d\'APIs tierces et développement d\'APIs RESTful')
    
    doc.add_heading('Formation', 1)
    doc.add_paragraph('Ingénieur en Génie Logiciel - ENSIAS Rabat')
    doc.add_paragraph('2017 - 2021')
    
    doc.add_heading('Compétences Techniques', 1)
    
    doc.add_heading('Stack Principal (Expert)', 2)
    doc.add_paragraph('• React: Hooks, Context API, Redux, React Router, Next.js')
    doc.add_paragraph('• Node.js: Express, NestJS, Socket.io, JWT, Passport')
    doc.add_paragraph('• Python: FastAPI, Django, Flask, Pandas, NumPy')
    doc.add_paragraph('• MongoDB: Mongoose, Aggregation, Indexing, Sharding')
    doc.add_paragraph('• Docker: Dockerfile, Docker Compose, Multi-stage builds, Volumes')
    
    doc.add_heading('Technologies Complémentaires', 2)
    doc.add_paragraph('• Frontend: TypeScript, JavaScript ES6+, HTML5, CSS3, Tailwind CSS')
    doc.add_paragraph('• Backend: RESTful APIs, GraphQL, Microservices')
    doc.add_paragraph('• Bases de données: MongoDB, PostgreSQL, Redis')
    doc.add_paragraph('• DevOps: Docker, Git, GitHub Actions, CI/CD')
    doc.add_paragraph('• Testing: Jest, Pytest, Mocha, Chai')
    doc.add_paragraph('• Outils: VS Code, Postman, MongoDB Compass, Docker Desktop')
    
    doc.add_heading('Projets Réalisés', 1)
    
    doc.add_paragraph('E-Commerce Platform (React + Node.js + MongoDB + Docker)')
    doc.add_paragraph('• Frontend React avec Redux pour la gestion d\'état')
    doc.add_paragraph('• Backend Node.js avec Express et MongoDB')
    doc.add_paragraph('• Containerisation complète avec Docker Compose')
    doc.add_paragraph('• Authentification JWT et paiement Stripe')
    
    doc.add_paragraph('Task Management System (React + Python + MongoDB + Docker)')
    doc.add_paragraph('• Interface React moderne et responsive')
    doc.add_paragraph('• API Python FastAPI avec MongoDB')
    doc.add_paragraph('• Déploiement avec Docker et Docker Compose')
    doc.add_paragraph('• WebSocket pour les notifications en temps réel')
    
    doc.add_paragraph('Analytics Dashboard (React + Node.js + MongoDB + Docker)')
    doc.add_paragraph('• Visualisation de données avec React et Recharts')
    doc.add_paragraph('• Backend Node.js avec agrégations MongoDB complexes')
    doc.add_paragraph('• Scripts Python pour le traitement de données')
    doc.add_paragraph('• Architecture microservices avec Docker')
    
    doc.add_heading('Certifications', 1)
    doc.add_paragraph('• MongoDB Certified Developer Associate')
    doc.add_paragraph('• Docker Certified Associate')
    doc.add_paragraph('• React - The Complete Guide (Udemy)')
    doc.add_paragraph('• Node.js - Advanced Concepts (Udemy)')
    
    doc.add_heading('Contributions Open Source', 1)
    doc.add_paragraph('• Contributeur actif sur GitHub (500+ contributions)')
    doc.add_paragraph('• Projets React, Node.js et Python')
    doc.add_paragraph('• Documentation et tutoriels Docker')
    
    doc.add_heading('Langues', 1)
    doc.add_paragraph('• Arabe: Langue maternelle')
    doc.add_paragraph('• Français: Courant')
    doc.add_paragraph('• Anglais: Avancé (technique)')
    
    doc.add_heading('Soft Skills', 1)
    doc.add_paragraph('• Travail en équipe Agile/Scrum')
    doc.add_paragraph('• Communication technique efficace')
    doc.add_paragraph('• Résolution de problèmes complexes')
    doc.add_paragraph('• Apprentissage continu et veille technologique')
    
    doc.save('cv_sarah_alami_fullstack.docx')
    print("✅ CV créé: cv_sarah_alami_fullstack.docx")
    print("\n📊 Compétences correspondantes:")
    print("   ✓ React")
    print("   ✓ Node.js")
    print("   ✓ Python")
    print("   ✓ MongoDB")
    print("   ✓ Docker")
    print("\n🎯 Ce CV devrait avoir un score de matching très élevé!")

if __name__ == "__main__":
    print("🚀 Génération du CV parfait pour l'offre Full Stack...\n")
    create_perfect_match_cv()
